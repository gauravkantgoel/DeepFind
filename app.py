"""
File Finder - Phase 2 Backend (Fixed)
pip install flask flask-cors watchdog sentence-transformers pymupdf python-docx openpyxl numpy
python app.py  →  open http://localhost:5000/ui

Fixes applied:
  - Stop Indexing now actually works (checked in both loops, reset on start)
  - Semantic AI tab populated independently (not gated behind keyword results)
  - Config file handle leak fixed
  - Config cached for watchdog (not re-read on every FS event)
  - Delta indexing: skip files whose mtime hasn't changed
  - Folder column stores actual parent directory (not watch root)
  - Thread-safe cache rebuild with proper locking
  - Path validation on open/open-folder endpoints
  - .doc extraction warning (python-docx only supports .docx)
"""

from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
import sqlite3, os, subprocess, sys, threading, time, json
from datetime import datetime
from pathlib import Path
import numpy as np

app = Flask(__name__)
CORS(app)

DB_PATH     = "file_index.db"
CONFIG_PATH = "config.json"

# ── Default config ─────────────────────────────────────────────────────────────

CODE_EXTENSIONS = [
    ".js",".jsx",".ts",".tsx",".py",".pyc",".pyo",".pyd",
    ".java",".class",".jar",".c",".cpp",".h",".hpp",".cs",
    ".go",".rs",".rb",".php",".swift",".kt",".scala",".sh",
    ".bat",".ps1",".cmd",".vbs",".r",".m",".lua",".pl",
    ".sql",".html",".htm",".css",".scss",".sass",".less",
    ".json",".xml",".yaml",".yml",".toml",".conf",".cfg",
    ".map",".min.js",".min.css",".lock",".gitignore",
    ".env",".bak",".swp",".sln",".csproj",".vcxproj",
    ".gradle",".cmake",".makefile",".dockerfile",
    ".whl",".egg",".gem",".so",".dylib",".o",".obj",
    ".node_modules",".DS_Store",
]

DEFAULT_CONFIG = {
    "watch_folders": [],
    "excluded_extensions": [".tmp",".log",".sys",".dll",".exe",
                            ".ini",".dat",".db",".lnk",".url",".pst",".ost"],
    "extraction": {
        "enabled":            True,
        "extract_types":      [".pdf",".docx",".xlsx",".xls",".txt",".csv"],
        "max_file_size_mb":   10,
        "content_length_cap": 2000,
        "throttle_ms":        100
    },
    "semantic_search": {
        "enabled":              True,
        "keyword_fallback":     True,
        "similarity_threshold": 0.40
    }
}

def get_effective_exclusions(cfg):
    """Exclusions for INDEXING only — system junk files, not code files."""
    return set(cfg.get("excluded_extensions", []))

def load_config():
    if os.path.exists(CONFIG_PATH):
        with open(CONFIG_PATH, "r") as f:          # FIX: proper file handle
            saved = json.load(f)
        for k, v in DEFAULT_CONFIG.items():
            if k not in saved:
                saved[k] = v
            elif isinstance(v, dict):
                for kk, vv in v.items():
                    if kk not in saved[k]:
                        saved[k][kk] = vv
        return saved
    return dict(DEFAULT_CONFIG)

def save_config(cfg):
    global _config_cache, _config_cache_time
    with open(CONFIG_PATH, "w") as f:
        json.dump(cfg, f, indent=2)
    # Invalidate cache on save
    _config_cache = cfg
    _config_cache_time = time.time()

# ── Cached config for hot paths (watchdog etc.) ──────────────────────────────

_config_cache      = None
_config_cache_time = 0
_CONFIG_TTL        = 5          # seconds

def get_cached_config():
    """Return config from cache if fresh, otherwise reload. Used by watchdog."""
    global _config_cache, _config_cache_time
    now = time.time()
    if _config_cache is not None and (now - _config_cache_time) < _CONFIG_TTL:
        return _config_cache
    _config_cache = load_config()
    _config_cache_time = now
    return _config_cache

# ── Database ───────────────────────────────────────────────────────────────────

_db_lock = threading.Lock()

def get_db():
    conn = sqlite3.connect(DB_PATH, timeout=30, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA busy_timeout=30000")
    return conn

def init_db():
    conn = get_db()
    conn.execute("""
        CREATE TABLE IF NOT EXISTS files (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            name        TEXT NOT NULL,
            path        TEXT NOT NULL UNIQUE,
            folder      TEXT NOT NULL,
            watch_root  TEXT,
            extension   TEXT,
            size        INTEGER,
            modified    REAL,
            indexed_at  REAL,
            content     TEXT,
            embedding   TEXT
        )
    """)
    for col, typ in [("content","TEXT"),("embedding","TEXT"),("watch_root","TEXT")]:
        try:
            conn.execute(f"ALTER TABLE files ADD COLUMN {col} {typ}")
        except Exception:
            pass
    for idx in [
        "CREATE INDEX IF NOT EXISTS idx_name       ON files(name)",
        "CREATE INDEX IF NOT EXISTS idx_folder     ON files(folder)",
        "CREATE INDEX IF NOT EXISTS idx_watchroot  ON files(watch_root)",
        "CREATE INDEX IF NOT EXISTS idx_ext        ON files(extension)",
        "CREATE INDEX IF NOT EXISTS idx_mod        ON files(modified DESC)",
    ]:
        conn.execute(idx)
    conn.commit()
    conn.close()

# ── Semantic model ─────────────────────────────────────────────────────────────

_model       = None
_model_lock  = threading.Lock()
model_status = {"loaded": False, "loading": False, "error": None}

def load_model():
    global _model, model_status
    cfg        = load_config()
    model_name = cfg.get("semantic_search", {}).get("model_name", "neuml/pubmedbert-base-embeddings")
    model_status.update({"loading": True, "error": None, "model_name": model_name})
    try:
        from sentence_transformers import SentenceTransformer
        m = SentenceTransformer(model_name)
        m.encode("warmup", convert_to_numpy=True)
        with _model_lock:
            _model = m
        model_status.update({"loaded": True, "loading": False})
        threading.Thread(target=rebuild_cache, daemon=True).start()
    except Exception as e:
        model_status.update({"error": str(e), "loaded": False, "loading": False})

def get_model():
    return _model

# ── Embedding cache ────────────────────────────────────────────────────────────

_emb_matrix  = None
_emb_paths   = []
_emb_lock    = threading.Lock()
_cache_state = {"ready": False, "building": False, "count": 0}

_rebuild_timer      = None
_rebuild_guard_lock = threading.Lock()       # FIX: proper guard for rebuild

def rebuild_cache():
    global _emb_matrix, _emb_paths

    if not model_status["loaded"]:
        return
    if indexing_status.get("running"):
        return

    # FIX: atomic check-and-set with lock to prevent double-builds
    with _rebuild_guard_lock:
        if _cache_state["building"]:
            return
        _cache_state.update({"building": True, "ready": False})

    try:
        conn = get_db()
        rows = conn.execute(
            "SELECT path, embedding FROM files WHERE embedding IS NOT NULL"
        ).fetchall()
        conn.close()

        if not rows:
            with _emb_lock:
                _emb_matrix = None
                _emb_paths  = []
            _cache_state.update({"count": 0})
            return

        paths, vecs = [], []
        for r in rows:
            try:
                v = json.loads(r["embedding"])
                paths.append(r["path"])
                vecs.append(v)
            except Exception:
                continue

        arr   = np.array(vecs, dtype=np.float32)
        norms = np.linalg.norm(arr, axis=1, keepdims=True)
        norms[norms == 0] = 1
        arr   = arr / norms

        with _emb_lock:
            _emb_matrix = arr
            _emb_paths  = paths
        _cache_state.update({"count": len(paths)})
    except Exception as e:
        print(f"Cache build error: {e}")
    finally:
        _cache_state.update({"building": False, "ready": True})

def schedule_cache_rebuild(delay=30):
    """Debounced rebuild — resets timer on every new event."""
    global _rebuild_timer
    if _rebuild_timer:
        _rebuild_timer.cancel()
    _rebuild_timer = threading.Timer(delay, rebuild_cache)
    _rebuild_timer.daemon = True
    _rebuild_timer.start()

# ── Pharma Synonym Dictionary ──────────────────────────────────────────────────

PHARMA_SYNONYMS = {
    "psur":  ["periodic safety update report","periodic benefit risk evaluation","pbrer"],
    "pbrer": ["periodic benefit risk evaluation report","psur","periodic safety update"],
    "dsur":  ["development safety update report","annual safety report"],
    "icsr":  ["individual case safety report","adverse event report","spontaneous report","case report"],
    "susar": ["suspected unexpected serious adverse reaction","unexpected serious adverse reaction"],
    "adr":   ["adverse drug reaction","adverse reaction","side effect","undesirable effect"],
    "ae":    ["adverse event","adverse experience","side effect"],
    "sae":   ["serious adverse event","serious adverse experience"],
    "cioms": ["cioms form","medwatch","individual case report form"],
    "rmp":   ["risk management plan","risk minimisation","risk minimization","safety specification"],
    "rmm":   ["risk minimisation measure","additional risk minimisation"],
    "rems":  ["risk evaluation and mitigation strategy","risk mitigation"],
    "dhcp":  ["dear healthcare professional","dear doctor letter","urgent safety restriction"],
    "signal detection": ["signal management","safety signal","signal assessment","disproportionality analysis"],
    "prr":   ["proportional reporting ratio","disproportionality","reporting ratio"],
    "ror":   ["reporting odds ratio","disproportionality analysis"],
    "ebgm":  ["empirical bayes geometric mean","bayesian signal detection"],
    "smpc":  ["summary of product characteristics","product information","prescribing information","label","spc"],
    "pil":   ["patient information leaflet","package leaflet"],
    "nda":   ["new drug application","marketing authorisation application","maa","bla"],
    "cta":   ["clinical trial application","investigational new drug","ind"],
    "epar":  ["european public assessment report","chmp opinion","assessment report"],
    "pharmacovigilance": ["drug safety","post marketing surveillance","pms","pv"],
    "case processing": ["icsr processing","case intake","narrative writing","case triage"],
    "literature review": ["literature search","published case reports","literature monitoring"],
    "aggregate report": ["periodic report","psur","pbrer","dsur","safety summary"],
    "benefit risk": ["benefit-risk evaluation","benefit risk assessment","pbrer"],
    "ema":   ["european medicines agency","chmp","european regulatory"],
    "fda":   ["food and drug administration","cder","us regulatory"],
    "mhra":  ["medicines and healthcare products regulatory agency","uk regulatory"],
    "pmda":  ["pharmaceuticals and medical devices agency","japan regulatory"],
    "who":   ["world health organization","vigibase","umc"],
    "meddra": ["medical dictionary for regulatory activities","preferred term","system organ class"],
    "soc":   ["system organ class","meddra soc"],
    "coding": ["medical coding","meddra coding","term selection"],
    "inspection": ["gvp inspection","regulatory inspection","audit","inspection readiness","mock inspection"],
    "gvp":   ["good pharmacovigilance practice","pharmacovigilance guidelines","gvp module"],
    "sop":   ["standard operating procedure","work instruction","process document"],
    "capa":  ["corrective action preventive action","corrective action","preventive action"],
    "audit": ["internal audit","external audit","compliance review","inspection"],
    "argus": ["oracle argus","argus safety","safety database"],
    "eudravigilance": ["ev","evweb","european adverse event database"],
    "vigibase": ["who vigibase","who database","umc database"],
    "faers": ["fda adverse event reporting system","fda spontaneous database"],
    "sdea":  ["safety data exchange agreement","pv agreement","pharmacovigilance agreement"],
    "pv agreement": ["safety data exchange agreement","sdea","pharmacovigilance contract"],
    "inn":   ["international nonproprietary name","generic name","active substance"],
    "protocol": ["study protocol","clinical protocol","trial protocol"],
    "training": ["training material","training record","e-learning"],
}

def expand_pharma_query(query):
    """
    Only expand multi-word queries or known full abbreviations.
    Skip expansion for very short single tokens to avoid false matches.
    """
    q_lower = query.lower().strip()
    words   = q_lower.split()

    is_multi_word   = len(words) > 1
    is_known_abbrev = any(q_lower == key for key in PHARMA_SYNONYMS)

    if not is_multi_word and not is_known_abbrev:
        return query, False

    extra = set()
    for key, synonyms in PHARMA_SYNONYMS.items():
        if key in q_lower or any(s in q_lower for s in synonyms):
            extra.update(synonyms)

    if not extra:
        return query, False

    new_terms = [t for t in extra if t.lower() not in q_lower][:4]
    if not new_terms:
        return query, False

    return (query + " " + " ".join(new_terms)).strip(), True

# ── Text extraction ────────────────────────────────────────────────────────────

def extract_text(path, ext, cfg):
    cap = cfg.get("extraction", {}).get("content_length_cap", 2000)
    try:
        if ext == ".pdf":
            import fitz
            doc  = fitz.open(path)
            text = " ".join(page.get_text() for page in doc[:5])
            doc.close()
        elif ext == ".docx":
            from docx import Document
            text = " ".join(p.text for p in Document(path).paragraphs if p.text.strip())
        elif ext == ".doc":
            # FIX: python-docx doesn't support .doc — try antiword, else skip
            try:
                result = subprocess.run(["antiword", path], capture_output=True, text=True, timeout=10)
                text = result.stdout if result.returncode == 0 else ""
            except (FileNotFoundError, subprocess.TimeoutExpired):
                print(f"  ⚠ Skipping .doc (antiword not installed): {path}")
                return None
        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb    = openpyxl.load_workbook(path, read_only=True, data_only=True)
            parts = []
            for ws in list(wb.worksheets)[:3]:
                for row in ws.iter_rows(max_row=100, values_only=True):
                    parts.extend(str(c) for c in row if c is not None and str(c).strip())
            text = " ".join(parts)
            wb.close()
        elif ext in (".txt", ".csv"):
            with open(path, encoding="utf-8", errors="ignore") as f:
                text = f.read()
        else:
            return None
        return text[:cap].strip() or None
    except Exception:
        return None

# ── Indexing ───────────────────────────────────────────────────────────────────

indexing_status = {
    "running": False, "count": 0, "current": "",
    "error": None, "embedded": 0, "skipped": 0,
    "stop_requested": False                         # FIX: explicit field
}

def index_folder(folder_path, tracker):
    cfg       = load_config()
    excl      = get_effective_exclusions(cfg)
    exc_cfg   = cfg.get("extraction", {})
    do_ext    = exc_cfg.get("enabled", True)
    ext_types = set(exc_cfg.get("extract_types", []))
    max_mb    = exc_cfg.get("max_file_size_mb", 10) * 1024 * 1024
    throttle  = exc_cfg.get("throttle_ms", 100) / 1000.0
    do_embed  = cfg.get("semantic_search", {}).get("enabled", True)
    model     = get_model() if do_embed else None
    conn      = get_db()

    for root, dirs, files in os.walk(folder_path):
        # ── FIX: check stop flag at directory level ──
        if tracker.get("stop_requested"):
            break

        dirs[:] = [d for d in dirs if not d.startswith(".")]
        for fname in files:
            # ── FIX: check stop flag at file level ──
            if tracker.get("stop_requested"):
                break

            if fname.startswith("."): continue
            ext   = Path(fname).suffix.lower()
            if ext in excl: continue
            fpath = os.path.join(root, fname)
            try:
                st = os.stat(fpath)

                # ── FIX: delta indexing — skip unchanged files ──
                existing = conn.execute(
                    "SELECT modified FROM files WHERE path = ?", (fpath,)
                ).fetchone()
                if existing and abs((existing["modified"] or 0) - st.st_mtime) < 0.001:
                    tracker["skipped"] = tracker.get("skipped", 0) + 1
                    continue

                content   = None
                embedding = None
                if do_ext and ext in ext_types and st.st_size <= max_mb:
                    content = extract_text(fpath, ext, cfg)
                if do_embed and model:
                    emb_text  = f"{fname} {content or ''}"
                    emb       = model.encode(emb_text, convert_to_numpy=True).tolist()
                    embedding = json.dumps(emb)
                    tracker["embedded"] += 1

                # Store both: watch_root (configured root) and folder (actual parent dir)
                actual_folder = os.path.dirname(fpath)

                conn.execute("""
                    INSERT OR REPLACE INTO files
                    (name, path, folder, watch_root, extension, size, modified, indexed_at, content, embedding)
                    VALUES (?,?,?,?,?,?,?,?,?,?)
                """, (fname, fpath, actual_folder, folder_path, ext,
                      st.st_size, st.st_mtime, time.time(), content, embedding))
                conn.commit()
                tracker["count"]   += 1
                tracker["current"]  = fpath
                if throttle > 0:
                    time.sleep(throttle)
            except (PermissionError, OSError):
                continue
    conn.close()

def run_indexing():
    global indexing_status
    cfg     = load_config()
    folders = cfg.get("watch_folders", [])

    # FIX: reset all status fields including stop_requested
    indexing_status.update({
        "running": True, "count": 0, "current": "",
        "error": None, "embedded": 0, "skipped": 0,
        "stop_requested": False
    })
    try:
        for folder in folders:
            # ── FIX: check stop flag between folders ──
            if indexing_status.get("stop_requested"):
                indexing_status["error"] = "Stopped by user"
                break

            if os.path.isdir(folder):
                index_folder(folder, indexing_status)
            else:
                indexing_status["error"] = f"Not found: {folder}"
    except Exception as e:
        indexing_status["error"] = str(e)
    finally:
        indexing_status["running"] = False
        threading.Thread(target=rebuild_cache, daemon=True).start()

# ── Watchdog ───────────────────────────────────────────────────────────────────

from watchdog.observers import Observer
from watchdog.events    import FileSystemEventHandler

class FileChangeHandler(FileSystemEventHandler):
    def __init__(self, root):
        self.root = root

    def _valid(self, path):
        cfg  = get_cached_config()
        excl = get_effective_exclusions(cfg)
        name = os.path.basename(path)
        return not name.startswith(".") and Path(name).suffix.lower() not in excl

    def on_created(self, event):
        if not event.is_directory and self._valid(event.src_path):
            self._upsert(event.src_path)

    def on_modified(self, event):
        if not event.is_directory and self._valid(event.src_path):
            self._upsert(event.src_path)

    def on_deleted(self, event):
        if not event.is_directory:
            try:
                conn = get_db()
                conn.execute("DELETE FROM files WHERE path = ?", (event.src_path,))
                conn.commit()
                conn.close()
            except Exception:
                pass
            schedule_cache_rebuild(30)

    def on_moved(self, event):
        if not event.is_directory:
            try:
                conn = get_db()
                conn.execute("DELETE FROM files WHERE path = ?", (event.src_path,))
                conn.commit()
                conn.close()
            except Exception:
                pass
            if self._valid(event.dest_path):
                self._upsert(event.dest_path)

    def _upsert(self, path):
        cfg     = get_cached_config()                # FIX: cached config
        exc_cfg = cfg.get("extraction", {})
        try:
            st    = os.stat(path)
            fname = os.path.basename(path)
            ext   = Path(fname).suffix.lower()
            max_mb= exc_cfg.get("max_file_size_mb", 10) * 1024 * 1024
            content = None; embedding = None
            if (exc_cfg.get("enabled", True)
                    and ext in set(exc_cfg.get("extract_types", []))
                    and st.st_size <= max_mb):
                content = extract_text(path, ext, cfg)
            model = get_model()
            if model and cfg.get("semantic_search", {}).get("enabled", True):
                emb       = model.encode(f"{fname} {content or ''}", convert_to_numpy=True).tolist()
                embedding = json.dumps(emb)

            actual_folder = os.path.dirname(path)
            watch_root    = self.root                # the configured watch folder

            conn = get_db()
            conn.execute("""
                INSERT OR REPLACE INTO files
                (name, path, folder, watch_root, extension, size, modified, indexed_at, content, embedding)
                VALUES (?,?,?,?,?,?,?,?,?,?)
            """, (fname, path, actual_folder, watch_root, ext,
                  st.st_size, st.st_mtime, time.time(), content, embedding))
            conn.commit(); conn.close()
            schedule_cache_rebuild(30)
        except (PermissionError, OSError):
            pass

observer     = None
watch_status = {"active": False, "folders": []}

def start_watching():
    global observer, watch_status
    cfg     = load_config()
    folders = [f for f in cfg.get("watch_folders", []) if os.path.isdir(f)]
    if not folders: return
    if observer and observer.is_alive():
        observer.stop(); observer.join()
    observer = Observer()
    for f in folders:
        observer.schedule(FileChangeHandler(f), f, recursive=True)
    observer.start()
    watch_status = {"active": True, "folders": folders}

# ── Search — 4-tier pipeline ───────────────────────────────────────────────────

def fmt_row(r, tier=1, score=None):
    return {
        "name":         r["name"],
        "path":         r["path"],
        "folder":       r["folder"],
        "extension":    r["extension"] or "",
        "size":         r["size"] or 0,
        "modified":     r["modified"],
        "modified_str": datetime.fromtimestamp(r["modified"]).strftime("%d %b %Y  %H:%M")
                        if r["modified"] else "—",
        "has_content":  bool(r["content"] if "content" in r.keys() else False),
        "score":        round(score * 100) if score is not None else None,
        "tier":         tier
    }

def build_where(extra_conds, extra_params, types, folder, exclude_paths=None):
    """Build SQL WHERE clause. folder uses path prefix matching."""
    conds  = list(extra_conds)
    params = list(extra_params)
    if types:
        exts = [f".{t.strip().lower()}" for t in types.split(",") if t.strip()]
        conds.append(f"extension IN ({','.join(['?']*len(exts))})")
        params.extend(exts)
    if folder:
        # Path prefix match — works for both watch roots and subfolders
        normalized = folder.replace("\\", "/").rstrip("/")
        conds.append("(REPLACE(path, '\\', '/') LIKE ?)")
        params.append(normalized + "/%")
    if exclude_paths:
        ep = list(exclude_paths)
        conds.append(f"path NOT IN ({','.join(['?']*len(ep))})")
        params.extend(ep)
    return " AND ".join(conds) if conds else "1=1", params

def tier1_all_keywords_in_filename(terms, types, folder, limit, seen):
    """All query terms present in filename."""
    name_conds = [f"LOWER(name) LIKE ?" for _ in terms]
    params     = [f"%{t}%" for t in terms]
    where, params = build_where(name_conds, params, types, folder, seen)
    conn = get_db()
    rows = conn.execute(
        f"SELECT * FROM files WHERE {where} ORDER BY modified DESC LIMIT ?",
        params + [limit]
    ).fetchall()
    conn.close()
    return [fmt_row(r, tier=1) for r in rows]

def tier2_some_keywords_in_filename(terms, types, folder, limit, seen):
    """At least one (but not all) query terms in filename."""
    if len(terms) == 1:
        return []
    or_conds = " OR ".join([f"LOWER(name) LIKE ?" for _ in terms])
    all_conds = " AND ".join([f"LOWER(name) LIKE ?" for _ in terms])
    conds  = [f"({or_conds})", f"NOT ({all_conds})"]
    params = [f"%{t}%" for t in terms] + [f"%{t}%" for t in terms]
    where, params = build_where(conds, params, types, folder, seen)
    conn = get_db()
    rows = conn.execute(
        f"SELECT * FROM files WHERE {where} ORDER BY modified DESC LIMIT ?",
        params + [limit]
    ).fetchall()
    conn.close()
    return [fmt_row(r, tier=2) for r in rows]

def tier3_content_keyword(terms, types, folder, limit, seen):
    """Any query term found in extracted document content (not filename)."""
    or_conds  = " OR ".join([f"LOWER(COALESCE(content,\'\')) LIKE ?" for _ in terms])
    name_none = " AND ".join([f"LOWER(name) NOT LIKE ?" for _ in terms])
    conds  = [f"({or_conds})", f"content IS NOT NULL", f"({name_none})"]
    params = [f"%{t}%" for t in terms] + [f"%{t}%" for t in terms]
    where, params = build_where(conds, params, types, folder, seen)
    conn = get_db()
    rows = conn.execute(
        f"SELECT * FROM files WHERE {where} ORDER BY modified DESC LIMIT ?",
        params + [limit]
    ).fetchall()
    conn.close()
    return [fmt_row(r, tier=3) for r in rows]

def tier4_semantic(query, types, folder, limit, threshold, keyword_paths):
    """
    Semantic similarity search.
    Returns (unique_results, overlap_count).
    """
    if not model_status["loaded"] or _cache_state["building"]:
        return [], 0
    with _emb_lock:
        matrix = _emb_matrix
        paths  = list(_emb_paths)
    if matrix is None or len(paths) == 0:
        return [], 0
    try:
        model = get_model()
        q_vec = model.encode(query, convert_to_numpy=True).astype(np.float32)
        q_vec = q_vec / (np.linalg.norm(q_vec) or 1)
        sims  = matrix.dot(q_vec)
    except Exception:
        return [], 0

    kw_set  = set(keyword_paths)
    unique  = []
    overlap = 0
    for i in range(len(paths)):
        if sims[i] >= threshold:
            if paths[i] in kw_set:
                overlap += 1
            else:
                unique.append((i, float(sims[i])))

    if not unique:
        return [], overlap

    unique.sort(key=lambda x: x[1], reverse=True)
    unique = unique[:limit]

    matching_paths = [paths[i] for i, _ in unique]
    path_to_score  = {paths[i]: s for i, s in unique}

    conds, params = [], []
    if types:
        exts = [f".{t.strip().lower()}" for t in types.split(",") if t.strip()]
        conds.append(f"extension IN ({','.join(['?']*len(exts))})")
        params.extend(exts)
    if folder:
        normalized = folder.replace("\\", "/").rstrip("/")
        conds.append("(REPLACE(path, '\\', '/') LIKE ?)")
        params.append(normalized + "/%")
    ph = ",".join(["?" for _ in matching_paths])
    conds.append(f"path IN ({ph})")
    params.extend(matching_paths)
    where = " AND ".join(conds) if conds else "1=1"

    conn = get_db()
    rows = conn.execute(f"SELECT * FROM files WHERE {where}", params).fetchall()
    conn.close()

    results = []
    for r in rows:
        score = path_to_score.get(r["path"], 0)
        results.append((score, fmt_row(r, tier=4, score=score)))
    results.sort(key=lambda x: x[0], reverse=True)
    return [r for _, r in results], overlap

_CODE_EXT_SET = set(CODE_EXTENSIONS)

def hybrid_search(query, types, folder, limit, threshold):
    """
    4-tier search pipeline.
    Returns (results, mode, semantic_overlap).
    """
    terms = [t for t in query.lower().split() if len(t) > 1]
    if not terms:
        return [], "keyword", 0

    seen    = set()
    results = []

    t1 = tier1_all_keywords_in_filename(terms, types, folder, limit, seen)
    results.extend(t1)
    seen.update(r["path"] for r in t1)

    remaining = limit - len(results)
    if remaining > 0:
        t2 = tier2_some_keywords_in_filename(terms, types, folder, remaining, seen)
        results.extend(t2)
        seen.update(r["path"] for r in t2)

    remaining = limit - len(results)
    if remaining > 0:
        t3 = tier3_content_keyword(terms, types, folder, remaining, seen)
        results.extend(t3)
        seen.update(r["path"] for r in t3)

    semantic_overlap = 0
    if model_status["loaded"] and not _cache_state["building"]:
        expanded_q, _ = expand_pharma_query(query)
        t4, semantic_overlap = tier4_semantic(expanded_q, types, folder, limit, threshold, seen)
        results.extend(t4)

    has_semantic = any(r["tier"] == 4 for r in results)
    warming_up   = model_status["loading"] or _cache_state["building"]
    mode = "semantic" if has_semantic else ("warming_up" if warming_up else "keyword")
    return results, mode, semantic_overlap

# ── API routes ─────────────────────────────────────────────────────────────────

@app.route("/api/status")
def api_status():
    conn  = get_db()
    total = conn.execute("SELECT COUNT(*) AS c FROM files").fetchone()["c"]
    emb_c = conn.execute("SELECT COUNT(*) AS c FROM files WHERE embedding IS NOT NULL").fetchone()["c"]
    conn.close()
    return jsonify({
        "indexed_files":  total,
        "embedded_files": emb_c,
        "indexing":       indexing_status,
        "watching":       watch_status,
        "model":          model_status,
        "cache":          _cache_state
    })

@app.route("/api/config", methods=["GET"])
def api_get_config():
    return jsonify(load_config())

@app.route("/api/config", methods=["POST"])
def api_save_config():
    save_config(request.json); return jsonify({"success": True})

@app.route("/api/index/stop", methods=["POST"])
def api_index_stop():
    if not indexing_status.get("running"):
        return jsonify({"error": "No indexing in progress"}), 400
    indexing_status["stop_requested"] = True
    return jsonify({"success": True, "message": "Stop signal sent"})

@app.route("/api/index", methods=["POST"])
def api_index():
    if indexing_status.get("running"):
        return jsonify({"error": "Indexing already running"}), 409
    if not load_config().get("watch_folders"):
        return jsonify({"error": "No folders configured"}), 400
    threading.Thread(target=run_indexing, daemon=True).start()
    return jsonify({"success": True})

@app.route("/api/clear", methods=["POST"])
def api_clear():
    conn = get_db()
    conn.execute("DELETE FROM files"); conn.commit(); conn.close()
    threading.Thread(target=rebuild_cache, daemon=True).start()
    return jsonify({"success": True})

@app.route("/api/clean", methods=["POST"])
def api_clean():
    conn  = get_db()
    rows  = conn.execute("SELECT path FROM files").fetchall()
    stale = [r["path"] for r in rows if not os.path.exists(r["path"])]
    if stale:
        conn.executemany("DELETE FROM files WHERE path = ?", [(p,) for p in stale])
        conn.commit()
    conn.close()
    if stale:
        threading.Thread(target=rebuild_cache, daemon=True).start()
    return jsonify({"removed": len(stale)})

@app.route("/api/search")
def api_search():
    query     = request.args.get("q", "").strip()
    types     = request.args.get("types", "").strip()
    folder    = request.args.get("folder", "").strip()
    hide_code = request.args.get("hide_code", "0") == "1"
    limit     = min(int(request.args.get("limit", 100)), 500)
    if not query: return jsonify({"results": [], "semantic_overlap": 0})

    cfg       = load_config()
    sem_cfg   = cfg.get("semantic_search", {})
    threshold = sem_cfg.get("similarity_threshold", 0.40)

    results, mode, semantic_overlap = hybrid_search(query, types, folder, limit, threshold)

    # Code file filtering — done in Python, not SQL (avoids 50+ param bloat)
    if hide_code:
        results = [r for r in results if r.get("extension","").lower() not in _CODE_EXT_SET]

    pharma_active = any(r.get("tier") == 4 for r in results)
    for r in results:
        r["search_mode"]     = mode
        r["pharma_expanded"] = pharma_active and mode == "semantic"

    return jsonify({
        "results":          results,
        "semantic_overlap":  semantic_overlap,
        "mode":             mode
    })

@app.route("/api/folders")
def api_folders():
    """Return configured watch folders with file counts."""
    cfg     = load_config()
    folders = cfg.get("watch_folders", [])
    conn    = get_db()

    result = []
    for f in folders:
        # Count files whose path starts with this watch folder
        normalized = f.replace("\\", "/").rstrip("/")
        cnt = conn.execute(
            "SELECT COUNT(*) AS c FROM files WHERE REPLACE(path, '\\', '/') LIKE ?",
            (normalized + "/%",)
        ).fetchone()["c"]
        name = os.path.basename(f.rstrip("/\\")) or f
        result.append({"path": f, "name": name, "count": cnt})

    conn.close()
    return jsonify(result)

def _validate_path(path):
    """FIX: ensure requested path is within an indexed watch folder."""
    cfg = load_config()
    watch = cfg.get("watch_folders", [])
    real = os.path.realpath(path)
    return any(real.startswith(os.path.realpath(w)) for w in watch)

@app.route("/api/open", methods=["POST"])
def api_open():
    path = request.json.get("path", "")
    if not os.path.exists(path):
        return jsonify({"error": "File not found"}), 404
    if not _validate_path(path):                     # FIX: path validation
        return jsonify({"error": "Path not in indexed folders"}), 403
    try:
        if sys.platform == "win32":    os.startfile(path)
        elif sys.platform == "darwin": subprocess.run(["open", path])
        else:                          subprocess.run(["xdg-open", path])
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/open-folder", methods=["POST"])
def api_open_folder():
    path   = request.json.get("path", "")
    folder = os.path.dirname(path)
    if not os.path.exists(folder):
        return jsonify({"error": "Folder not found"}), 404
    if not _validate_path(path):                     # FIX: path validation
        return jsonify({"error": "Path not in indexed folders"}), 403
    try:
        if sys.platform == "win32":
            subprocess.Popen(f'explorer /select,"{path}"')
        elif sys.platform == "darwin":
            subprocess.run(["open", "-R", path])
        else:
            subprocess.run(["xdg-open", folder])
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/ui")
def ui(): return send_file("index.html")

# ── Entry point ────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    init_db()
    threading.Thread(target=load_model,     daemon=True).start()
    threading.Thread(target=start_watching, daemon=True).start()
    def open_browser():
        time.sleep(1.5)
        import webbrowser
        webbrowser.open("http://localhost:5000/ui")
    threading.Thread(target=open_browser, daemon=True).start()
    print("\n" + "="*52)
    print("  File Finder Phase 2 is running!")
    print("  Open http://localhost:5000/ui")
    print("="*52 + "\n")
    app.run(debug=False, port=5000)
